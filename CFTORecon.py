"""Compare a supplied WCS CSV or WMS workbook with yesterday's Sterling releases.

Configure the variables below, then run this standalone script.
Requires paramiko locally, and openpyxl for Excel only.
The jump server needs Python 3 and oracledb. No SQLPlus or SSH tunnel is used.
"""

import argparse
import csv
from datetime import datetime
from html import escape
from itertools import islice
import json
from pathlib import Path
import re
import shlex
import socket
import sys


SCRIPT_DIR = Path(__file__).resolve().parent

STERLING_DB_USER = ""
STERLING_DB_PASSWORD = ""

JUMP_HOST = "YOUR_JUMP_HOST"
JUMP_PORT = 22
JUMP_USER = "YOUR_JUMP_USER"
KEY_PATH = "PATH_TO_SSH_PRIVATE_KEY"
JUMP_PYTHON = "python3"
KNOWN_HOSTS = SCRIPT_DIR / "known_hosts"
SSH_CONNECT_TIMEOUT = 20
QUERY_TIMEOUT = 900

STERLING_DSN = ("(DESCRIPTION =(ADDRESS = (PROTOCOL = TCP)(HOST = YOUR_DB_HOST)(PORT = 1521))"
                "(CONNECT_DATA =(SERVICE_NAME = YOUR_DB_SERVICE)))")

STERLING_RELEASE_QUERY = """select distinct trim(ms.key2) from STRADMIN1.ms_msg_replay ms, STRADMIN1.SPKV_YFS_ORDER_HEADER YOH, STRADMIN1.SPKV_YFS_ORDER_LINE YOL where
YOH.ORDER_HEADER_KEY=YOL.ORDER_HEADER_KEY AND trim(ms.key2)=trim(YOL.EXTN_FP_RELEASE_ID) AND to_char(ms.createts, 'DD-MON-YYYY') = to_char(sysdate-1, 'DD-MON-YYYY') and ms.Key2 <> '0' and
ms.msg_type_destination = 'DN_FOM' and YOH.ORDER_TYPE IN('CFTO_UK','CFTO_IE') and yoh.enterprise_key = 'MS_UK' and yoh.order_closed in ('N')  and yol.ordered_qty > '0'"""

REMOTE_QUERY_SCRIPT = """
import json
import sys


def main():
    settings = {}
    try:
        settings = json.load(sys.stdin)
        try:
            import oracledb
        except ImportError as error:
            raise RuntimeError("Python package 'oracledb' is unavailable on the jump server. "
                               "Set JUMP_PYTHON to a Python 3 environment with it installed.") from error

        with oracledb.connect(
            user=settings["user"], password=settings["password"], dsn=settings["dsn"],
            tcp_connect_timeout=settings["connect_timeout"], retry_count=0, disable_oob=True,
        ) as connection:
            connection.call_timeout = settings["query_timeout"] * 1000
            with connection.cursor() as cursor:
                cursor.execute(settings["query"].rstrip().rstrip(";"))
                orders = [row[0] for row in cursor if row[0] is not None]
        print(json.dumps({"orders": orders}))
        return 0
    except Exception as error:
        message = str(error)
        password = settings.get("password", "")
        if password:
            message = message.replace(password, "[REDACTED]")
        print(json.dumps({"error": "Jump-server Oracle query failed: " + message}))
        return 1


sys.exit(main())
"""

WMS_EXCEL_PATH = SCRIPT_DIR / "WCS_CFTO_Daily_reconciliation-2026-10-26.csv"
WMS_SHEET_NAME = None
WMS_HEADER_ROW = 1
WMS_ORDER_COLUMN = "ORDER_NUMBER"
OUTPUT_HTML = SCRIPT_DIR / "CFTOReconReport.html"

def open_jump_socket():
    try:
        return socket.create_connection((JUMP_HOST, JUMP_PORT), timeout=SSH_CONNECT_TIMEOUT)
    except socket.gaierror as error:
        raise RuntimeError("Cannot resolve SSH jump host '{}'. Check corporate VPN and DNS.".format(
            JUMP_HOST)) from error
    except socket.timeout as error:
        raise RuntimeError("SSH TCP connection to {}:{} timed out after {} seconds. "
                           "Check corporate VPN, firewall access and the SSH port; "
                           "the Oracle database has not been contacted.".format(
                               JUMP_HOST, JUMP_PORT, SSH_CONNECT_TIMEOUT)) from error
    except OSError as error:
        raise RuntimeError("Cannot open SSH TCP connection to {}:{}: {}. "
                           "Confirm the jump host, SSH port and network access.".format(
                               JUMP_HOST, JUMP_PORT, error)) from error


def normalise_id(value, number_format=""):
    if value is None:
        return ""
    if isinstance(value, str):
        return value.strip()
    if isinstance(value, bool):
        raise ValueError("An order ID cannot be a Boolean value.")
    if isinstance(value, int) or isinstance(value, float) and value.is_integer():
        identifier = str(int(value))
        if re.fullmatch(r"0+", number_format) and value >= 0:
            identifier = identifier.zfill(len(number_format))
        return identifier
    raise ValueError("Order IDs must be text or whole numbers.")


def read_wms_orders(path, sheet_name, header_row, order_column):
    path = Path(path)
    if not path.is_file():
        raise ValueError("Input file not found: {}".format(path))
    if path.suffix.lower() not in (".csv", ".xlsx", ".xlsm"):
        raise ValueError("Supply a .csv, .xlsx or .xlsm input file.")
    if header_row < 1:
        raise ValueError("WMS_HEADER_ROW must be at least 1.")

    if path.suffix.lower() == ".csv":
        with path.open(encoding="utf-8-sig", newline="") as source:
            reader = csv.reader(source, strict=True)
            header = next(islice(reader, header_row - 1, header_row), [])
            headers = [name.strip() for name in header]
            matches = [index for index, name in enumerate(headers)
                       if name.casefold() == order_column.strip().casefold()]
            if len(matches) != 1:
                raise ValueError("Expected one '{}' column in CSV; found headers: {}".format(
                    order_column, ", ".join(headers)))

            order_index = matches[0]
            orders = []
            for row_number, values in enumerate(reader, header_row + 1):
                if not any(value.strip() for value in values):
                    continue
                if order_index >= len(values):
                    raise ValueError("CSV row {} has no order number.".format(row_number))
                identifier = normalise_id(values[order_index])
                if not identifier:
                    raise ValueError("CSV row {} has a blank order number.".format(row_number))
                orders.append(identifier)
            if not orders:
                raise ValueError("The CSV contains no order numbers; no report was generated.")
            return orders, "CSV"

    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet_name is not None and sheet_name not in workbook.sheetnames:
            raise ValueError("WMS worksheet not found: {}".format(sheet_name))
        worksheet = workbook[sheet_name] if sheet_name is not None else workbook.worksheets[0]
        header = next(worksheet.iter_rows(min_row=header_row, max_row=header_row), ())
        headers = [str(cell.value).strip() if cell.value is not None else "" for cell in header]
        matches = [index for index, name in enumerate(headers)
                   if name.casefold() == order_column.strip().casefold()]
        if len(matches) != 1:
            raise ValueError("Expected one '{}' column in '{}'; found headers: {}".format(
                order_column, worksheet.title, ", ".join(headers)))

        order_index = matches[0]
        orders = []
        for row_number, cells in enumerate(worksheet.iter_rows(min_row=header_row + 1), header_row + 1):
            if not any(cell.value is not None and str(cell.value).strip() for cell in cells):
                continue
            if order_index >= len(cells):
                raise ValueError("WMS row {} has no order ID.".format(row_number))
            cell = cells[order_index]
            try:
                identifier = normalise_id(cell.value, cell.number_format)
            except ValueError as error:
                raise ValueError("WMS row {}: {}".format(row_number, error)) from error
            if not identifier:
                raise ValueError("WMS row {} has a blank order ID.".format(row_number))
            orders.append(identifier)
        if not orders:
            raise ValueError("The WMS worksheet contains no order IDs; no report was generated.")
        return orders, worksheet.title
    finally:
        workbook.close()


def fetch_sterling_orders():
    if not STERLING_DB_USER.strip() or not STERLING_DB_PASSWORD:
        raise ValueError("Set STERLING_DB_USER and STERLING_DB_PASSWORD in CFTORecon.py before running.")
    if not Path(KEY_PATH).is_file():
        raise ValueError("SSH private key not found: {}. Set KEY_PATH for this computer.".format(
            KEY_PATH))

    import paramiko

    client = paramiko.SSHClient()
    connection_socket = None
    try:
        connection_socket = open_jump_socket()
        client.load_system_host_keys()
        Path(KNOWN_HOSTS).touch(exist_ok=True)
        client.load_host_keys(str(KNOWN_HOSTS))
        client.set_missing_host_key_policy(paramiko.AutoAddPolicy())
        try:
            client.connect(hostname=JUMP_HOST, port=JUMP_PORT, username=JUMP_USER,
                           key_filename=KEY_PATH, look_for_keys=False, allow_agent=False,
                           sock=connection_socket, timeout=SSH_CONNECT_TIMEOUT,
                           auth_timeout=20, banner_timeout=20)
        except paramiko.AuthenticationException as error:
            raise RuntimeError("SSH authentication failed. Check JUMP_USER and KEY_PATH; "
                               "the Oracle database has not been contacted.") from error
        except socket.timeout as error:
            raise RuntimeError("SSH handshake or authentication timed out. TCP connected, "
                               "but the jump server did not complete SSH login.") from error
        except paramiko.SSHException as error:
            raise RuntimeError("SSH login failed: {}".format(error)) from error

        transport = client.get_transport()
        if transport is None or not transport.is_active():
            raise RuntimeError("The SSH connection closed before the remote query started.")
        transport.set_keepalive(30)

        command = "{} -c {}".format(shlex.quote(JUMP_PYTHON), shlex.quote(REMOTE_QUERY_SCRIPT))
        stdin, stdout, stderr = client.exec_command(
            command, timeout=QUERY_TIMEOUT + SSH_CONNECT_TIMEOUT)
        stdin.write(json.dumps({
            "user": STERLING_DB_USER,
            "password": STERLING_DB_PASSWORD,
            "dsn": STERLING_DSN,
            "query": STERLING_RELEASE_QUERY,
            "connect_timeout": SSH_CONNECT_TIMEOUT,
            "query_timeout": QUERY_TIMEOUT,
        }))
        stdin.flush()
        stdin.channel.shutdown_write()
        try:
            output = stdout.read().decode("utf-8", errors="replace")
            error_output = stderr.read().decode("utf-8", errors="replace")
            status = stdout.channel.recv_exit_status()
        except socket.timeout as error:
            raise RuntimeError("Timed out waiting for the Oracle query on the jump server.") from error

        try:
            result = json.loads(output)
        except json.JSONDecodeError as error:
            detail = error_output.strip().replace(STERLING_DB_PASSWORD, "[REDACTED]")
            raise RuntimeError("Jump-server Python returned no valid query result (exit {}). "
                               "Check JUMP_PYTHON and its Python 3 environment. {}".format(
                                   status, detail)) from error
        if not isinstance(result, dict):
            raise RuntimeError("Jump-server Python returned an invalid query result.")
        if status != 0 or result.get("error"):
            message = str(result.get("error") or "Remote Python exited with status {}.".format(status))
            raise RuntimeError(message.replace(STERLING_DB_PASSWORD, "[REDACTED]"))
        if not isinstance(result.get("orders"), list):
            raise RuntimeError("Jump-server Python did not return an order-ID list.")
        orders = [normalise_id(value) for value in result["orders"]]
        return [identifier for identifier in orders if identifier]
    finally:
        client.close()
        if connection_socket is not None:
            connection_socket.close()


def compare_orders(wms_orders, sterling_orders):
    wms_unique = list(dict.fromkeys(wms_orders))
    sterling_unique = set(sterling_orders)
    missing = [identifier for identifier in wms_unique if identifier not in sterling_unique]
    return {
        "wms_count": len(wms_unique),
        "sterling_count": len(sterling_unique),
        "matched_count": len(wms_unique) - len(missing),
        "missing": missing,
        "duplicate_wms_count": len(wms_orders) - len(wms_unique),
    }


def build_html_report(result, wms_path, sheet_name):
    missing_count = len(result["missing"])
    status = "GREEN" if missing_count == 0 else "AMBER" if missing_count < 50 else "RED"
    generated = datetime.now().astimezone().strftime("%d %b %Y %H:%M:%S %Z")
    source_details = ("Format: CSV" if Path(wms_path).suffix.lower() == ".csv"
                      else "Worksheet: {}".format(sheet_name))
    rows = "".join("<tr><td>{}</td></tr>".format(escape(identifier))
                   for identifier in result["missing"])
    missing_section = (
        '<h2>Orders Missing In Sterling</h2><table><thead><tr><th scope="col">{}</th>'
        '</tr></thead><tbody>{}</tbody></table>'.format(escape(WMS_ORDER_COLUMN), rows)
        if missing_count else '<p class="matched">All input order IDs were found in the Sterling extract.</p>'
    )
    return """<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>CFTO Reconciliation Report</title>
<style>
:root {{ --ink: #20262b; --muted: #56616b; --line: #d6dde2; --paper: #fff; --band: #edf3f5; }}
* {{ box-sizing: border-box; }}
body {{ margin: 0; background: #f6f8f9; color: var(--ink); font: 16px/1.5 Calibri, Candara, sans-serif; letter-spacing: 0; }}
main {{ max-width: 1000px; margin: 0 auto; padding: 28px 20px 48px; background: var(--paper); min-height: 100vh; }}
header {{ border-bottom: 3px solid #217176; padding-bottom: 16px; }}
h1 {{ font-size: 26px; margin: 0 0 8px; }}
h2 {{ font-size: 20px; margin-top: 28px; }}
p {{ margin: 8px 0; overflow-wrap: anywhere; }}
.metadata {{ color: var(--muted); font-size: 15px; }}
.status {{ font-weight: bold; padding: 10px 12px; margin-top: 20px; border-left: 4px solid; }}
.GREEN, .matched {{ color: #185b32; background: #edf7f0; }}
.AMBER {{ color: #714600; background: #fff5d9; }}
.RED {{ color: #9c2026; background: #fff0f0; }}
.matched {{ padding: 12px; }}
table {{ border-collapse: collapse; width: 100%; margin-top: 16px; table-layout: fixed; }}
th, td {{ text-align: left; padding: 10px 12px; border: 1px solid var(--line); overflow-wrap: anywhere; vertical-align: top; }}
th {{ background: var(--band); }}
.summary th:last-child, .summary td:last-child {{ width: 100px; text-align: right; }}
@media (max-width: 480px) {{ main {{ padding: 18px 12px 32px; }} h1 {{ font-size: 23px; }} th, td {{ padding: 8px; }} }}
@media print {{ body {{ background: white; }} main {{ max-width: none; min-height: 0; padding: 0; }} tr {{ break-inside: avoid; }} }}
</style>
</head>
<body><main>
<header><h1>CFTO Reconciliation Report</h1>
<p class="metadata">Generated: {generated}</p>
<p class="metadata">Source file: {filename} | {source_details}</p>
<p class="metadata">Period: yesterday, according to the Sterling database date (SYSDATE).</p></header>
<p class="status {status}">Reconciliation status: {status}</p>
<table class="summary"><thead><tr><th scope="col">Category</th><th scope="col">Count</th></tr></thead><tbody>
<tr><td>Unique input order IDs</td><td>{wms_count}</td></tr>
<tr><td>Unique Sterling release IDs</td><td>{sterling_count}</td></tr>
<tr><td>Input order IDs found in Sterling</td><td>{matched_count}</td></tr>
<tr><td>Input order IDs missing in Sterling</td><td>{missing_count}</td></tr>
<tr><td>Duplicate input rows excluded from counts</td><td>{duplicates}</td></tr>
</tbody></table>
{missing_section}
</main></body></html>
""".format(
    generated=escape(generated), filename=escape(Path(wms_path).name),
    source_details=escape(source_details),
        status=status, wms_count=result["wms_count"], sterling_count=result["sterling_count"],
        matched_count=result["matched_count"], missing_count=missing_count,
        duplicates=result["duplicate_wms_count"], missing_section=missing_section,
    )


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--check-ssh", action="store_true",
                        help="Test jump-host TCP connectivity without credentials or a database query.")
    arguments = parser.parse_args(argv)
    try:
        if arguments.check_ssh:
            connection_socket = open_jump_socket()
            connection_socket.close()
            print("PASS: TCP connection to {}:{} succeeded. "
                  "SSH authentication and the Oracle database were not checked.".format(
                      JUMP_HOST, JUMP_PORT))
            return 0
        wms_orders, sheet_name = read_wms_orders(
            WMS_EXCEL_PATH, WMS_SHEET_NAME, WMS_HEADER_ROW, WMS_ORDER_COLUMN)
        sterling_orders = fetch_sterling_orders()
        result = compare_orders(wms_orders, sterling_orders)
        report = build_html_report(result, WMS_EXCEL_PATH, sheet_name)
        OUTPUT_HTML.write_text(report, encoding="utf-8")
        print("Input: {} unique IDs | Matched: {} | Missing in Sterling: {}".format(
            result["wms_count"], result["matched_count"], len(result["missing"])))
        print("HTML report: {}".format(OUTPUT_HTML))
        return 0
    except Exception as error:
        message = str(error)
        if STERLING_DB_PASSWORD:
            message = message.replace(STERLING_DB_PASSWORD, "[REDACTED]")
        print("FAILED: {}".format(message), file=sys.stderr)
        return 1


if __name__ == "__main__":
    sys.exit(main())